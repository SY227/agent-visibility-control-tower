from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "agent_visibility_control_tower_pitch_deck.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(0xF6, 0xFB, 0xF7)
MIST = RGBColor(0xEE, 0xF6, 0xEF)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x15, 0x20, 0x16)
SLATE = RGBColor(0x49, 0x56, 0x4A)
MUTED = RGBColor(0x6F, 0x7D, 0x70)
BORDER = RGBColor(0xD8, 0xE4, 0xD9)
GREEN = RGBColor(0x3E, 0x8F, 0x5C)
GREEN_DEEP = RGBColor(0x2D, 0x6D, 0x45)
GREEN_SOFT = RGBColor(0xDF, 0xF2, 0xE5)
SAGE = RGBColor(0x8D, 0xBB, 0x95)
SAGE_SOFT = RGBColor(0xED, 0xF7, 0xEF)
AMBER = RGBColor(0xD1, 0xA5, 0x42)
AMBER_SOFT = RGBColor(0xFB, 0xF3, 0xDF)
RED = RGBColor(0xC9, 0x69, 0x5A)
RED_SOFT = RGBColor(0xFD, 0xEE, 0xEA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Aptos"
FONT_DISPLAY = "Aptos"


def set_background(slide, dark=False):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = GREEN_DEEP if dark else BG

    if not dark:
        add_oval(slide, 9.75, -0.15, 3.1, 2.4, GREEN_SOFT, transparency=0.15, line=False)
        add_oval(slide, -0.6, -0.2, 3.4, 2.1, SAGE, transparency=0.78, line=False)
        add_oval(slide, 8.55, 5.45, 3.2, 1.5, GREEN, transparency=0.90, line=False)


def add_oval(slide, x, y, w, h, color, transparency=0.0, line=False):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.fill.transparency = transparency
    if line:
        shape.line.color.rgb = color
    else:
        shape.line.fill.background()
    return shape


def add_box(slide, x, y, w, h, fill_color=CARD, line_color=BORDER, radius=True, transparency=0.0, line_width=1.0):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.fill.transparency = transparency
    shape.line.color.rgb = line_color
    shape.line.width = Pt(line_width)
    return shape


def add_text(slide, x, y, w, h, text, size=18, color=INK, bold=False, font=FONT, align=PP_ALIGN.LEFT,
             valign=MSO_ANCHOR.TOP, italic=False, margin=0.08, line_spacing=1.1):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin * 0.75)
    tf.margin_bottom = Inches(margin * 0.45)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def add_badge(slide, x, y, text, fill_color=GREEN_SOFT, text_color=GREEN_DEEP, w=None):
    width = w if w is not None else max(1.2, min(3.7, 0.075 * len(text) + 0.55))
    shape = add_box(slide, x, y, width, 0.34, fill_color=fill_color, line_color=fill_color, radius=True)
    shape.line.fill.background()
    add_text(slide, x + 0.04, y + 0.01, width - 0.08, 0.26, text, size=9.5, color=text_color, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    return shape


def add_title_block(slide, eyebrow, title, subtitle=None):
    add_badge(slide, 0.72, 0.46, eyebrow)
    add_text(slide, 0.72, 0.95, 11.1, 0.7, title, size=26, color=INK, bold=True, font=FONT_DISPLAY)
    if subtitle:
        add_text(slide, 0.72, 1.52, 11.2, 0.42, subtitle, size=12.5, color=SLATE)


def add_page_num(slide, n):
    add_text(slide, 12.52, 7.0, 0.35, 0.22, str(n), size=9.5, color=MUTED, align=PP_ALIGN.RIGHT)


def add_placeholder(slide, x, y, w, h, label):
    shape = add_box(slide, x, y, w, h, fill_color=WHITE, line_color=SAGE, radius=True, transparency=0.02)
    shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    add_text(slide, x + 0.12, y + 0.12, w - 0.24, h - 0.24, f"[INSERT APP SCREENSHOT: {label}]", size=13, color=GREEN_DEEP, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    return shape


def add_flow_box(slide, x, y, w, h, title, tone="green", subtitle=None):
    if tone == "green":
        fill_color, text_color, line_color = WHITE, INK, BORDER
    elif tone == "soft":
        fill_color, text_color, line_color = SAGE_SOFT, INK, SAGE
    elif tone == "amber":
        fill_color, text_color, line_color = AMBER_SOFT, INK, AMBER
    else:
        fill_color, text_color, line_color = GREEN_SOFT, GREEN_DEEP, GREEN
    add_box(slide, x, y, w, h, fill_color=fill_color, line_color=line_color)
    add_text(slide, x + 0.08, y + 0.1, w - 0.16, 0.34, title, size=12.5, color=text_color, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    if subtitle:
        add_text(slide, x + 0.12, y + 0.42, w - 0.24, h - 0.5, subtitle, size=9.5, color=SLATE, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def add_arrow_text(slide, x, y, text="→", size=18, color=MUTED):
    add_text(slide, x, y, 0.34, 0.28, text, size=size, color=color, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def add_bullet_card(slide, x, y, w, h, title, bullets, tone="green"):
    fill_color = WHITE
    line_color = BORDER
    badge_fill = GREEN_SOFT
    badge_text = GREEN_DEEP
    if tone == "amber":
        badge_fill, badge_text, line_color = AMBER_SOFT, RGBColor(0x8A, 0x69, 0x21), AMBER
    elif tone == "red":
        badge_fill, badge_text, line_color = RED_SOFT, RGBColor(0x9B, 0x4D, 0x41), RED
    elif tone == "sage":
        badge_fill, badge_text, line_color = SAGE_SOFT, GREEN_DEEP, SAGE
    add_box(slide, x, y, w, h, fill_color=fill_color, line_color=line_color)
    add_badge(slide, x + 0.14, y + 0.12, title, fill_color=badge_fill, text_color=badge_text, w=min(w - 0.28, max(1.5, 0.075 * len(title) + 0.55)))
    cur_y = y + 0.56
    for bullet in bullets:
        add_text(slide, x + 0.18, cur_y, w - 0.36, 0.42, f"• {bullet}", size=10.5, color=SLATE)
        cur_y += 0.38


def add_stat_card(slide, x, y, w, h, title, value, detail=None, fill_color=WHITE, line_color=BORDER, value_color=INK):
    add_box(slide, x, y, w, h, fill_color=fill_color, line_color=line_color)
    add_text(slide, x + 0.14, y + 0.12, w - 0.28, 0.24, title, size=10, color=MUTED, bold=True)
    add_text(slide, x + 0.14, y + 0.43, w - 0.28, 0.38, value, size=16, color=value_color, bold=True)
    if detail:
        add_text(slide, x + 0.14, y + 0.82, w - 0.28, h - 0.92, detail, size=10.2, color=SLATE)


def add_chip_row(slide, items, x, y, max_w=12.0):
    cur_x = x
    cur_y = y
    for item in items:
        width = max(1.3, min(2.7, 0.078 * len(item) + 0.55))
        if cur_x + width > x + max_w:
            cur_x = x
            cur_y += 0.38
        add_badge(slide, cur_x, cur_y, item, fill_color=SAGE_SOFT, text_color=GREEN_DEEP, w=width)
        cur_x += width + 0.12


# Slide 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_badge(slide, 0.72, 0.5, "Transforming Enterprise Through AI | Track 2", w=2.95)
add_text(slide, 0.72, 1.1, 6.2, 1.15, "Agent Visibility\nControl Tower", size=28, color=INK, bold=True, font=FONT_DISPLAY, line_spacing=0.95)
add_text(slide, 0.72, 2.55, 5.8, 0.38, "The control tower for the machine-facing GTM layer", size=15, color=GREEN_DEEP, bold=True)
add_text(slide, 0.72, 3.0, 5.95, 0.9, "Inspect and repair how AI agents, LLMs, and answer engines understand, cite, route, or skip your company.", size=13.2, color=SLATE)
add_text(slide, 0.72, 6.45, 6.0, 0.3, "One URL in. One bounded Gemini workflow. One executive-ready Fix Pack.", size=10.5, color=MUTED)
hero = add_box(slide, 7.55, 1.0, 5.0, 5.45, fill_color=WHITE, line_color=BORDER)
add_text(slide, 7.88, 1.25, 4.34, 0.32, "Boardroom-ready executive work product", size=11, color=GREEN_DEEP, bold=True, align=PP_ALIGN.CENTER)
add_placeholder(slide, 7.88, 1.72, 4.34, 3.5, "Hero / main app screen")
add_box(slide, 7.88, 5.42, 4.34, 0.66, fill_color=GREEN_SOFT, line_color=GREEN_SOFT)
add_text(slide, 8.04, 5.56, 4.0, 0.22, "Premium, calm, evidence-constrained workflow for enterprise teams", size=10.3, color=GREEN_DEEP, align=PP_ALIGN.CENTER, bold=True)
add_page_num(slide, 1)

# Slide 2
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title_block(slide, "Market shift", "Your next visitor isn’t human.", "A machine can now shape perception, routing, and trust before a human buyer ever clicks.")
add_text(slide, 0.88, 2.1, 1.05, 0.22, "Old web", size=11, color=MUTED, bold=True)
old_y = 2.45
old_xs = [0.88, 2.42, 3.96, 5.5]
old_labels = ["Search", "Click", "Website", "Human\nevaluates"]
for i, (x, label) in enumerate(zip(old_xs, old_labels)):
    add_flow_box(slide, x, old_y, 1.18, 0.78, label, tone="green")
    if i < len(old_xs) - 1:
        add_arrow_text(slide, x + 1.23, old_y + 0.23)
add_text(slide, 0.88, 4.1, 1.15, 0.22, "New web", size=11, color=MUTED, bold=True)
new_boxes = [
    (0.88, 4.45, 1.7, 0.88, "LLM / AI agent"),
    (3.0, 4.45, 3.05, 0.88, "Summarizes / cites / routes / skips"),
    (6.45, 4.45, 2.2, 0.88, "Buyer may\nnever click"),
]
for i, (x, y, w, h, label) in enumerate(new_boxes):
    add_flow_box(slide, x, y, w, h, label, tone="soft")
    if i < len(new_boxes) - 1:
        add_arrow_text(slide, x + w + 0.14, y + 0.27)
callout = add_box(slide, 9.2, 2.25, 3.25, 3.6, fill_color=WHITE, line_color=BORDER)
add_badge(slide, 9.42, 2.48, "What changed", fill_color=AMBER_SOFT, text_color=RGBColor(0x8A, 0x69, 0x21), w=1.3)
add_text(slide, 9.42, 2.9, 2.82, 0.85, "A new machine-facing GTM layer now sits before the website experience.", size=14.5, color=INK, bold=True)
for idx, txt in enumerate([
    "AI-generated summaries can frame the company before the buyer visits the site.",
    "Answer engines can cite, compare, route, or skip a company with no click at all.",
    "Most enterprises do not yet inspect or manage that layer directly.",
]):
    add_text(slide, 9.42, 3.88 + idx * 0.45, 2.82, 0.35, f"• {txt}", size=10.4, color=SLATE)
add_box(slide, 0.88, 6.28, 11.58, 0.55, fill_color=GREEN_SOFT, line_color=GREEN_SOFT)
add_text(slide, 1.0, 6.43, 11.3, 0.22, "If AI systems shape perception before the click, visibility risk moves upstream.", size=11.5, color=GREEN_DEEP, bold=True, align=PP_ALIGN.CENTER)
add_page_num(slide, 2)

# Slide 3
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title_block(slide, "Problem", "Most enterprise websites were built for humans, not machine-mediated buyers.", "The result is not just discoverability risk, but interpretation, citation, and routing risk.")
add_bullet_card(slide, 0.88, 2.1, 2.82, 1.65, "Misclassification", [
    "Category framing is often implicit.",
    "Models can flatten the company into a broader label.",
], tone="red")
add_bullet_card(slide, 3.9, 2.1, 2.82, 1.65, "Citation weakness", [
    "Proof exists, but is not packaged for reuse.",
    "Claims may be paraphrased without strong support.",
], tone="amber")
add_bullet_card(slide, 6.92, 2.1, 2.82, 1.65, "Agent journey blocker", [
    "Buyer routing is fragmented.",
    "The next best action is not always machine-clear.",
], tone="sage")
add_bullet_card(slide, 9.94, 2.1, 2.5, 1.65, "Buyer routing risk", [
    "Machines may miss the right path.",
    "Good products can still be skipped.",
], tone="green")
add_box(slide, 0.88, 4.28, 11.56, 1.58, fill_color=WHITE, line_color=BORDER)
add_text(slide, 1.1, 4.55, 11.1, 0.36, "What enterprises risk", size=11, color=MUTED, bold=True)
add_text(slide, 1.1, 4.9, 11.0, 0.62, "A company can be directionally legible to humans and still weakly interpreted by AI systems that summarize, compare, and route buyer intent.", size=18, color=INK, bold=True)
add_box(slide, 0.88, 6.15, 11.56, 0.56, fill_color=RED_SOFT, line_color=RED_SOFT)
add_text(slide, 1.0, 6.3, 11.25, 0.22, "This is not a ranking problem alone. It is an interpretation, citation, and routing problem.", size=11.4, color=RGBColor(0x9B, 0x4D, 0x41), bold=True, align=PP_ALIGN.CENTER)
add_page_num(slide, 3)

# Slide 4
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title_block(slide, "Solution", "One URL in. A boardroom-ready AI Visibility Brief and Fix Pack out.", "A bounded Gemini workflow with visible artifacts and evidence-constrained outputs, not a fake autonomous swarm.")
add_flow_box(slide, 0.85, 2.4, 1.28, 0.86, "Website URL", tone="feature", subtitle="Single input")
add_arrow_text(slide, 2.22, 2.65)
add_flow_box(slide, 2.58, 2.12, 1.48, 0.96, "Website\nContext Agent", subtitle="Site facts captured")
add_arrow_text(slide, 4.13, 2.44)
add_flow_box(slide, 4.48, 2.12, 1.48, 0.96, "LLM\nPerception Agent", subtitle="Likely machine summary")
add_arrow_text(slide, 6.03, 2.44)
add_flow_box(slide, 6.38, 2.12, 1.48, 0.96, "Agent\nVisitor Agent", subtitle="Journey blockers")
add_arrow_text(slide, 7.93, 2.44)
add_flow_box(slide, 8.28, 2.12, 1.58, 0.96, "AIO / Answer\nEngine Agent", subtitle="Answer-surface findings")
add_arrow_text(slide, 9.97, 2.44)
add_flow_box(slide, 10.32, 2.12, 1.58, 0.96, "Citation\nReadiness Agent", subtitle="Evidence gaps")
add_arrow_text(slide, 11.99, 2.44)
add_flow_box(slide, 10.32, 3.74, 1.58, 0.96, "Fix\nPrioritization Agent", subtitle="Top actions + Fix Pack")
add_text(slide, 11.13, 3.18, 0.28, 0.42, "↓", size=18, color=MUTED, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
add_arrow_text(slide, 9.97, 4.05, text="←")
add_arrow_text(slide, 8.33, 4.05, text="←")
add_flow_box(slide, 6.58, 3.74, 2.95, 0.96, "AI Visibility Readiness Brief", tone="feature", subtitle="Executive Verdict | Boardroom Snapshot | Decision Memo | Fix Pack")
add_box(slide, 0.88, 5.52, 7.25, 0.96, fill_color=WHITE, line_color=BORDER)
add_badge(slide, 1.05, 5.78, "Workflow design principles", fill_color=GREEN_SOFT, text_color=GREEN_DEEP, w=1.78)
add_chip_row(slide, ["Bounded public-web scan", "Typed / inspectable artifacts", "Evidence-constrained outputs", "No fake autonomous swarm"], 1.02, 6.13, max_w=6.75)
add_placeholder(slide, 8.48, 5.52, 3.96, 0.96, "Workflow trace")
add_page_num(slide, 4)

# Slide 5
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title_block(slide, "Product output", "From scan to executive action plan in one workflow", "What the product actually produces is a real executive work product, not just a model response.")
add_chip_row(slide, [
    "Executive Verdict", "Boardroom Snapshot", "Decision Memo", "Scorecards",
    "Fix Pack", "Machine-Facing GTM Risks", "Evidence Receipts", "Visible Agent Artifacts",
], 0.88, 1.98, max_w=11.6)
add_placeholder(slide, 0.88, 2.7, 5.7, 2.1, "Executive Verdict + Boardroom Snapshot")
add_placeholder(slide, 0.88, 5.02, 2.74, 1.28, "Decision Memo + Scorecards")
add_placeholder(slide, 3.84, 5.02, 2.74, 1.28, "Fix Pack")
add_box(slide, 6.9, 2.7, 5.5, 3.6, fill_color=WHITE, line_color=BORDER)
add_badge(slide, 7.12, 2.95, "Why judges should believe it", fill_color=AMBER_SOFT, text_color=RGBColor(0x8A, 0x69, 0x21), w=1.92)
add_text(slide, 7.12, 3.38, 5.0, 0.45, "The output reads like something a GTM, content, or executive team can act on immediately.", size=14.5, color=INK, bold=True)
for idx, txt in enumerate([
    "It compresses risk, consequence, and first move into boardroom language.",
    "It exposes visible agent artifacts instead of hiding the workflow behind a black box.",
    "It hands teams an actionable Fix Pack with content, proof, routing, and schema workstreams.",
    "It keeps limitations explicit and avoids unverified claims or fake precision.",
]):
    add_text(slide, 7.12, 4.06 + idx * 0.43, 5.0, 0.34, f"• {txt}", size=10.4, color=SLATE)
add_page_num(slide, 5)

# Slide 6
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title_block(slide, "Shopify case study", "Demo case: Shopify", "The strongest current demo run shows why this matters even for a high-quality, machine-readable site.")
add_stat_card(slide, 0.88, 2.06, 2.05, 1.22, "AI Visibility Score", "96 / 100", "Excellent", fill_color=GREEN_SOFT, line_color=GREEN_SOFT, value_color=GREEN_DEEP)
add_stat_card(slide, 0.88, 3.5, 4.1, 1.25, "Top machine-facing gap", "Buyer path present, but pricing and evaluation logic may still be fragmented.", None, fill_color=WHITE, line_color=BORDER)
add_stat_card(slide, 0.88, 4.98, 4.1, 1.15, "First-week move", "Make the B2B, DTC, and enterprise distinction more machine-readable.", None, fill_color=AMBER_SOFT, line_color=AMBER, value_color=INK)
add_bullet_card(slide, 0.88, 6.2, 4.1, 0.95, "Example Fix Pack actions", [
    "Package proof closer to homepage and enterprise claims.",
    "Clarify the enterprise evaluation path.",
    "Reinforce product, organization, FAQ, and enterprise schema.",
], tone="green")
add_placeholder(slide, 5.32, 2.05, 3.25, 2.0, "Shopify Executive Verdict")
add_placeholder(slide, 8.84, 2.05, 3.25, 2.0, "Shopify Fix Pack")
add_placeholder(slide, 5.32, 4.38, 6.77, 1.78, "Shopify Evidence Receipts or Agent Artifacts")
add_box(slide, 5.32, 6.42, 6.77, 0.52, fill_color=GREEN_SOFT, line_color=GREEN_SOFT)
add_text(slide, 5.5, 6.57, 6.4, 0.18, "Strong sites still benefit when machine-facing GTM signals are easier to interpret, cite, and route.", size=10.8, color=GREEN_DEEP, bold=True, align=PP_ALIGN.CENTER)
add_page_num(slide, 6)

# Slide 7
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title_block(slide, "Why Gemini / Track 2", "Built for Track 2: Gemini-powered agent workflows for enterprise use", "Gemini is used for bounded orchestration, typed outputs, and visible intermediate artifacts that improve trust and usability.")
capabilities = [
    ("Bounded six-stage workflow", "Gemini powers the core workflow, stage by stage."),
    ("Structured, typed outputs", "Artifacts are constrained and render cleanly into an executive brief."),
    ("Visible intermediate artifacts", "Inspectability improves enterprise trust and usability."),
    ("Evidence-constrained synthesis", "Reliability is prioritized over fake autonomous behavior."),
]
for idx, (title, body) in enumerate(capabilities):
    x = 0.88 + idx * 3.0
    add_box(slide, x, 2.0, 2.72, 1.25, fill_color=WHITE, line_color=BORDER)
    add_text(slide, x + 0.14, 2.18, 2.45, 0.28, title, size=11.2, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, x + 0.14, 2.52, 2.45, 0.5, body, size=9.6, color=SLATE, align=PP_ALIGN.CENTER)
criteria = [
    ("Application of Technology", "Gemini drives the workflow, structured synthesis, visible artifacts, and Fix Pack generation."),
    ("Presentation", "One URL input, clear workflow trace, and boardroom-ready outputs make the demo legible fast."),
    ("Business Value", "Enterprises risk being misunderstood, weakly cited, or misrouted before a human buyer ever clicks."),
    ("Originality", "This audits and repairs the AI-era visibility layer, not generic SEO or crawler health."),
]
coords = [(0.88, 3.62), (6.7, 3.62), (0.88, 5.35), (6.7, 5.35)]
for (title, body), (x, y) in zip(criteria, coords):
    add_box(slide, x, y, 5.42, 1.42, fill_color=WHITE, line_color=BORDER)
    add_text(slide, x + 0.18, y + 0.16, 5.0, 0.26, title, size=11.2, color=GREEN_DEEP, bold=True)
    add_text(slide, x + 0.18, y + 0.48, 5.0, 0.72, body, size=10.2, color=SLATE)
add_page_num(slide, 7)

# Slide 8
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title_block(slide, "Enterprise relevance", "This is machine-facing GTM infrastructure, not a generic audit.", "Different enterprise teams can act on the same output without needing a redesign or a new data stack.")
lanes = [
    ("GTM / Growth", "Identify where AI systems may misread the company, flatten differentiation, or fail to route buyer intent."),
    ("Digital / Content", "Repair answer-engine readiness, proof packaging, and citation support on the highest-signal pages."),
    ("Engineering", "Implement schema, machine-readable structure, and page-level reinforcement without changing the product itself."),
    ("Executives", "See risk, consequence, ownership, and the first-week move in boardroom language."),
]
for idx, (title, body) in enumerate(lanes):
    x = 0.88 + (idx % 2) * 6.0
    y = 2.15 + (idx // 2) * 1.95
    fill_color = WHITE if idx % 2 == 0 else SAGE_SOFT
    line_color = BORDER if idx % 2 == 0 else SAGE
    add_box(slide, x, y, 5.56, 1.55, fill_color=fill_color, line_color=line_color)
    add_badge(slide, x + 0.16, y + 0.16, title, fill_color=GREEN_SOFT if idx % 2 == 0 else WHITE, text_color=GREEN_DEEP, w=1.4 if title == "Executives" else None)
    add_text(slide, x + 0.16, y + 0.58, 5.15, 0.66, body, size=11.4, color=SLATE)
add_box(slide, 0.88, 6.35, 11.56, 0.5, fill_color=AMBER_SOFT, line_color=AMBER_SOFT)
add_text(slide, 1.04, 6.49, 11.2, 0.18, "The output is designed to align strategy, content, and implementation around one machine-facing risk layer.", size=10.8, color=RGBColor(0x8A, 0x69, 0x21), bold=True, align=PP_ALIGN.CENTER)
add_page_num(slide, 8)

# Slide 9
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_badge(slide, 0.72, 0.5, "Closing")
add_text(slide, 0.72, 1.0, 6.4, 1.2, "Websites were built for humans.\nThe next layer must work for machines too.", size=27, color=INK, bold=True, font=FONT_DISPLAY, line_spacing=0.96)
add_text(slide, 0.72, 2.7, 5.95, 0.7, "Agent Visibility Control Tower gives enterprises a way to inspect and repair how AI systems perceive them before a buyer ever clicks.", size=13.4, color=SLATE)
add_chip_row(slide, ["Machine-facing GTM layer", "Bounded Gemini workflow", "Executive-ready Fix Pack"], 0.72, 3.72, max_w=5.6)
add_box(slide, 0.72, 4.35, 5.9, 1.25, fill_color=WHITE, line_color=BORDER)
add_text(slide, 0.95, 4.65, 5.42, 0.2, "Final footer line", size=10.2, color=MUTED, bold=True)
add_text(slide, 0.95, 4.98, 5.42, 0.32, "One URL. One bounded Gemini workflow. One executive-ready Fix Pack.", size=15, color=GREEN_DEEP, bold=True)
add_placeholder(slide, 7.15, 1.0, 5.05, 3.55, "Final app hero screenshot")
add_box(slide, 7.15, 4.88, 5.05, 1.48, fill_color=WHITE, line_color=BORDER)
add_text(slide, 7.35, 5.08, 4.7, 0.24, "Screenshots to capture and insert manually", size=11.2, color=GREEN_DEEP, bold=True)
manual = [
    "Hero / main app screen",
    "Workflow trace",
    "Executive Verdict + Boardroom Snapshot",
    "Decision Memo + Scorecards",
    "Fix Pack",
    "Shopify case output",
    "Evidence Receipts / Agent Artifacts",
]
for idx, item in enumerate(manual):
    add_text(slide, 7.35, 5.33 + idx * 0.145, 4.6, 0.14, f"• {item}", size=8.2, color=SLATE)
add_page_num(slide, 9)

prs.save(str(OUTPUT))
print(f"Wrote {OUTPUT}")
